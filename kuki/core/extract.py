from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx"}
SUPPORTED_EXTENSIONS = sorted(TEXT_EXTENSIONS | IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS)


@dataclass(frozen=True)
class OCRStatus:
    ready: bool
    backend: str
    message: str


def get_ocr_status() -> OCRStatus:
    try:
        from rapidocr_onnxruntime import RapidOCR  # noqa: F401
    except Exception as exc:
        return OCRStatus(
            ready=False,
            backend="RapidOCR",
            message=f"Image OCR is unavailable until rapidocr-onnxruntime is installed: {exc}",
        )

    return OCRStatus(
        ready=True,
        backend="RapidOCR",
        message="Image OCR is ready for local CPU extraction.",
    )


@lru_cache(maxsize=1)
def _get_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


def _extract_from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except Exception:
        return ""

    reader = PdfReader(str(path))
    chunks = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            chunks.append(text)
    return "\n".join(chunks)


def _extract_from_docx(path: Path) -> str:
    try:
        from docx import Document
    except Exception:
        return ""

    doc = Document(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())


def _extract_from_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except Exception:
        return ""

    presentation = Presentation(str(path))
    chunks = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = (shape.text or "").strip()
                if text:
                    chunks.append(text)
    return "\n".join(chunks)


def _extract_from_image(path: Path) -> str:
    status = get_ocr_status()
    if not status.ready:
        return ""

    result, _elapsed = _get_ocr_engine()(str(path))
    lines = []
    for item in result or []:
        if not isinstance(item, (list, tuple)) or len(item) < 2:
            continue
        raw_text = item[1]
        if isinstance(raw_text, str):
            text = raw_text
        elif isinstance(raw_text, (list, tuple)) and raw_text:
            text = str(raw_text[0])
        else:
            text = str(raw_text)
        text = text.strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def extract_text_from_file(path: Path) -> str:
    ext = path.suffix.lower()

    if ext in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        return _extract_from_pdf(path)
    if ext == ".docx":
        return _extract_from_docx(path)
    if ext == ".pptx":
        return _extract_from_pptx(path)
    if ext in IMAGE_EXTENSIONS:
        return _extract_from_image(path)

    return ""


def preview_text(text: str, max_chars: int = 1200) -> str:
    cleaned = " ".join((text or "").split())
    if len(cleaned) <= max_chars:
        return cleaned
    return f"{cleaned[:max_chars].rstrip()}..."
